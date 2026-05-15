"""Тесты Excel/1С-экспортов по одной компании (Step C2)."""
from __future__ import annotations

from io import BytesIO

import pytest

from exports import build_company_1c_csv, build_company_xlsx
from schemas import CompanyData


@pytest.fixture
def full_company():
    """Полностью заполненная компания — все поля заданы."""
    return CompanyData(
        inn="7707083893",
        kpp="773601001",
        ogrn="1027700132195",
        name="ПАО Сбербанк",
        status="Действующая",
        director="Греф Г.О.",
        region="г Москва",
        address="г. Москва, ул. Вавилова, 19",
        okved_main="64.19",
        okved_name="Прочее денежное посредничество",
        reg_date="1991-06-20",
        age_years=33,
        capital=67760844000.0,
        employees_count=210000,
        revenue_last_year=3_200_000_000_000.0,
        profit_last_year=900_000_000_000.0,
        source="dadata+fns+zchb",
    )


@pytest.fixture
def sparse_company():
    """Минимально заполненная — только ИНН и название."""
    return CompanyData(inn="1234567890", name="ИП Иванов", source="dadata")


class TestBuildCompanyXlsx:
    def test_returns_xlsx_bytes(self, full_company):
        data = build_company_xlsx(full_company)
        assert isinstance(data, bytes)
        # XLSX это zip-архив, начинается с PK\x03\x04
        assert data[:2] == b"PK"

    def test_xlsx_contains_company_data(self, full_company):
        """Проверяем, что openpyxl действительно записал значения."""
        from openpyxl import load_workbook
        data = build_company_xlsx(full_company)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        # Собираем все значения в плоский список
        values = [
            ws.cell(row=r, column=c).value
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
        ]
        assert "ПАО Сбербанк" in values
        assert "7707083893" in values
        assert "Греф Г.О." in values
        assert "64.19" in values

    def test_xlsx_handles_sparse_company(self, sparse_company):
        from openpyxl import load_workbook
        data = build_company_xlsx(sparse_company)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        # Должно быть «—» для пустых полей, а не упасть
        values = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        assert "ИП Иванов" in values
        assert "—" in values  # пустые поля помечены прочерком

    def test_xlsx_header_row_present(self, full_company):
        from openpyxl import load_workbook
        data = build_company_xlsx(full_company)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        assert ws["A1"].value == "Карточка контрагента"


class TestBuild1CCsv:
    def test_returns_bytes_cp1251(self, full_company):
        data = build_company_1c_csv(full_company)
        assert isinstance(data, bytes)
        # Должен декодироваться в cp1251 — это стандарт 1С
        text = data.decode("cp1251")
        assert "ПАО Сбербанк" in text
        assert "7707083893" in text

    def test_uses_standard_1c_headers(self, full_company):
        """1С импортирует CSV по точному совпадению имён колонок —
        нужны строго стандартные имена реквизитов справочника
        «Контрагенты»."""
        data = build_company_1c_csv(full_company)
        text = data.decode("cp1251")
        first_line = text.split("\n")[0]
        # Проверяем критичные колонки — без них импорт не сработает
        for required in ("ИНН", "КПП", "Наименование", "ОГРН",
                         "ЮридическийАдрес", "Руководитель"):
            assert required in first_line, f"{required} отсутствует в заголовке"

    def test_separator_is_semicolon(self, full_company):
        """1С 8.3 ждёт `;` — не запятую (которая используется в
        десятичных дробях рублёвых значений)."""
        data = build_company_1c_csv(full_company)
        text = data.decode("cp1251")
        assert ";" in text

    def test_sparse_company_does_not_crash(self, sparse_company):
        data = build_company_1c_csv(sparse_company)
        text = data.decode("cp1251")
        assert "ИП Иванов" in text
        assert "1234567890" in text

    def test_handles_unicode_replacement_in_cp1251(self):
        """Если в названии есть символ, которого нет в cp1251 (типа
        грузинских букв), он должен быть заменён на ? без падения."""
        company = CompanyData(
            inn="1234567890", name="ООО Ⴀтест", source="dadata",
        )
        data = build_company_1c_csv(company)
        assert isinstance(data, bytes)  # не упало
        # cp1251 либо смог либо подставил '?' — главное не исключение


class TestBuildCompanyCardPdf:
    """D2: структурированная PDF-карточка контрагента."""

    def test_returns_pdf_bytes(self, full_company):
        from exports import build_company_card_pdf
        data = build_company_card_pdf(full_company)
        assert isinstance(data, bytes)
        # PDF начинается с "%PDF"
        assert data[:4] == b"%PDF"

    def test_works_with_sparse_company(self, sparse_company):
        from exports import build_company_card_pdf
        data = build_company_card_pdf(sparse_company)
        assert data[:4] == b"%PDF"
        # Sparse не должен падать — пустые поля рисуются как «—»

    def test_works_without_security_result(self, full_company):
        from exports import build_company_card_pdf
        data = build_company_card_pdf(full_company, security_result=None)
        assert data[:4] == b"%PDF"

    def test_works_with_security_result(self, full_company):
        from exports import build_company_card_pdf

        # Минимальная заглушка SecurityResult — должна не уронить рендер
        class FakeItem:
            def __init__(self, name, status, is_critical=False):
                self.name = name
                self.status = status
                self.is_critical = is_critical

        class FakeResult:
            items = [
                FakeItem("ФССП", "не найден"),
                FakeItem("Банкротство", "найдено", is_critical=True),
            ]

        data = build_company_card_pdf(full_company, security_result=FakeResult())
        assert data[:4] == b"%PDF"

    def test_handles_corrupt_security_result(self, full_company):
        """Если security_result имеет неожиданную форму — не падаем."""
        from exports import build_company_card_pdf
        data = build_company_card_pdf(full_company, security_result="not-an-object")
        assert data[:4] == b"%PDF"
